import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def is_md_table_line(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and s.count("|") >= 2


def is_md_table_sep(line: str) -> bool:
    s = line.strip()
    if not is_md_table_line(s):
        return False
    cells = [c.strip() for c in s.strip("|").split("|")]
    return all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def parse_md_table(lines: list[str], start_idx: int) -> tuple[list[list[str]], int]:
    # returns (rows, next_index)
    header = [c.strip() for c in lines[start_idx].strip().strip("|").split("|")]
    i = start_idx + 1
    if i < len(lines) and is_md_table_sep(lines[i]):
        i += 1
    rows: list[list[str]] = [header]
    while i < len(lines) and is_md_table_line(lines[i]) and not is_md_table_sep(lines[i]):
        row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        rows.append(row)
        i += 1
    return rows, i


_INLINE_RE = re.compile(r"(\*\*(.+?)\*\*|`([^`]+)`)")

def add_inline_runs(paragraph, text: str) -> None:
    pos = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        if match.group(2) is not None:
            run = paragraph.add_run(match.group(2))
            run.bold = True
        else:
            run = paragraph.add_run(match.group(3))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_formatted_paragraph(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    return paragraph


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _set_cell_borders(cell, color: str = "CCCCCC") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), color)
        borders.append(border)
    tc_pr.append(borders)


def _set_cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for name, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        margin = OxmlElement(f"w:{name}")
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    tc_pr.append(margins)


def add_code_block(doc: Document, code_lines: list[str], language: str | None = None) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.text = ""

    _set_cell_shading(cell, "F5F5F5")
    _set_cell_borders(cell)
    _set_cell_margins(cell)

    if language:
        label = cell.paragraphs[0]
        label_run = label.add_run(language)
        label_run.font.name = "Calibri"
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        label_run.italic = True
        label.paragraph_format.space_after = Pt(4)
        label.paragraph_format.space_before = Pt(0)

    if not code_lines:
        code_lines = [""]

    for idx, code_line in enumerate(code_lines):
        paragraph = cell.paragraphs[0] if idx == 0 and not language else cell.add_paragraph()
        run = paragraph.add_run(code_line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0

    doc.add_paragraph("")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows) if rows else 0
    if cols == 0:
        return
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            add_inline_runs(cell.paragraphs[0], text)
    doc.add_paragraph("")


_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")


def add_heading_line(doc: Document, line: str):
    match = _HEADING_RE.match(line)
    if not match:
        return False
    level = len(match.group(1)) - 1
    title = match.group(2).strip()
    style = "Title" if level == 0 else f"Heading {level}"
    paragraph = doc.add_paragraph(style=style)
    if level == 0:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline_runs(paragraph, title)
    return True


def is_horizontal_rule(line: str) -> bool:
    return bool(re.fullmatch(r"\s*-{3,}\s*", line))


def add_horizontal_rule(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")
    p_borders.append(bottom)
    p_pr.append(p_borders)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)


def main() -> None:
    project_root = Path(__file__).resolve().parent
    if len(sys.argv) >= 2:
        md_path = Path(sys.argv[1])
        out_path = Path(sys.argv[2]) if len(sys.argv) >= 3 else md_path.with_suffix(".docx")
    else:
        md_path = project_root / "docs" / "NC2-2_test_report_2.md"
        out_path = project_root / "docs" / "NC2-2_test_report_2.docx"

    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()

    # default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        if not line.strip():
            i += 1
            continue

        # headings (# .. ######)
        if add_heading_line(doc, line):
            i += 1
            continue

        if is_horizontal_rule(line):
            add_horizontal_rule(doc)
            i += 1
            continue

        # table
        if is_md_table_line(line):
            rows, next_i = parse_md_table(lines, i)
            add_table(doc, rows)
            i = next_i
            continue

        # bullet list
        if re.match(r"^\s*-\s+", line):
            add_formatted_paragraph(doc, re.sub(r"^\s*-\s+", "", line).strip(), style="List Bullet")
            i += 1
            continue

        # numbered list
        if re.match(r"^\s*\d+\.\s+", line):
            add_formatted_paragraph(doc, re.sub(r"^\s*\d+\.\s+", "", line).strip(), style="List Number")
            i += 1
            continue

        # code fence
        if line.strip().startswith("```"):
            fence = line.strip()
            language = fence[3:].strip() or None
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            if i < len(lines) and lines[i].strip().startswith("```"):
                i += 1
            add_code_block(doc, code_lines, language)
            continue

        # regular paragraph with inline **bold** and `code`
        add_formatted_paragraph(doc, line)
        i += 1

    try:
        doc.save(out_path)
        print(f"Saved: {out_path}")
    except PermissionError:
        alt_path = out_path.with_name(out_path.stem + "_new.docx")
        try:
            doc.save(alt_path)
            print(f"Файл занят, сохранено в: {alt_path}")
        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            fallback_path = out_path.with_name(f"{out_path.stem}_{timestamp}.docx")
            doc.save(fallback_path)
            print(f"Файлы заняты, сохранено в: {fallback_path}")


if __name__ == "__main__":
    main()

